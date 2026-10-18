"""P2: paginación en servidor de /productos, /inventario, /clientes y /proveedores.

- 25 sintéticos por listado dentro de transacciones con rollback (tienda
  para productos/movimientos/proveedores; auth para clientes), sin datos
  personales.
- La página de cada fila se deriva de su POSICIÓN en el orden del
  listado, calculada en runtime: los tests no hardcodean los tamaños
  reales ni dependen de que los reales no cambien.
- ?pagina= inválida (abc, 0, -3, 1.5) → 1; mayor que la última → última
  con "Mostrando X–Y de Z" correcto, nunca 500.
- Totales y contadores sobre el conjunto COMPLETO: idénticos en todas
  las páginas.
- Las exportaciones no se pagan: el workbook lleva TODO el conjunto.
"""
import io
import re
from pathlib import Path

import psycopg2
import psycopg2.extras
import pytest
from openpyxl import load_workbook

import db_supabase
import app as modulo_app
from app import app

_TEMPLATES = Path(__file__).resolve().parent.parent / 'templates'
POR_PAGINA = 20
NSINT = 25   # sintéticos: llena una página y rueda el resto a la siguiente


def _admin():
    c = app.test_client()
    with c.session_transaction() as s:
        s['usuario_id'] = 7
        s['rol'] = 'administrador'
        s['nombre'] = 'Prueba'
        s['_csrf_token'] = 'token-test'
    return c


class _Proxy:
    """Contrato de conexión proxied con rollback (como P1)."""

    def __init__(self, raw):
        self._raw = raw

    def cursor(self, dictionary=False, **kwargs):
        if dictionary:
            return self._raw.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        return self._raw.cursor()

    def commit(self):
        pass

    def rollback(self):
        self._raw.rollback()

    def close(self):
        pass


def _pag_de(posicion):
    """Página (1-index) en la que cae la fila de la posición dada."""
    return (posicion + POR_PAGINA - 1) // POR_PAGINA


def _filas_excel(resp):
    wb = load_workbook(io.BytesIO(resp.get_data()))
    return [fila for fila in wb.active.iter_rows(values_only=True)]


@pytest.fixture()
def fx_tienda(monkeypatch):
    raw = db_supabase.engine_tienda.raw_connection()
    proxy = _Proxy(raw)
    cur = proxy.cursor(dictionary=True)
    # 25 productos sintéticos: por id quedan al FINAL del listado
    cur.execute(f"""
        INSERT INTO producto (nombre, precio, stock)
        SELECT 'testp2-prod-' || gs || '-' || md5(random()::text), 9.99, 50
          FROM generate_series(1, {NSINT}) AS gs
        RETURNING id_producto
    """)
    ids_prod = sorted(r['id_producto'] for r in cur.fetchall())
    # 25 movimientos sintéticos con fechas distintas: por fecha DESC
    # quedan al INICIO del historial
    cur.execute("""
        INSERT INTO inventario_movimiento
            (tipo, id_producto, cantidad, precio_unitario, observacion,
             id_usuario, stock_resultante, fecha)
        SELECT 'entrada', %s, 1, 9.99, 'testp2-mov-' || gs, 8, 50,
               NOW() - (gs || ' minutes')::interval
          FROM generate_series(1, %s) AS gs
    """, (ids_prod[0], NSINT))
    # 25 proveedores sintéticos 'zzz-': por nombre quedan al FINAL
    cur.execute("""
        INSERT INTO proveedor (nombre, contacto, telefono, email, direccion)
        SELECT 'zzz-testp2-prov-' || lpad(gs::text, 2, '0'), '', '',
               'testp2-' || gs || '@invalido', ''
          FROM generate_series(1, %s) AS gs
        RETURNING id_proveedor
    """, (NSINT,))
    totales = {}
    for tabla in ('producto', 'inventario_movimiento', 'proveedor'):
        cur.execute(f"SELECT COUNT(*) AS n FROM {tabla}")
        totales[tabla] = cur.fetchone()['n']
    # Categoría real con productos y sin espacios, para filtros en href
    cur.execute("""
        SELECT c.nombre, COUNT(*) AS n
          FROM categoria c
          JOIN producto p ON p.id_categoria = c.id_categoria
         WHERE c.nombre NOT LIKE '% %'
         GROUP BY c.nombre, c.id_categoria
         ORDER BY c.id_categoria LIMIT 1
    """)
    fila = cur.fetchone()
    categoria = fila['nombre'] if fila else None
    if categoria:
        totales['categoria'] = fila['n']
    monkeypatch.setattr(modulo_app, 'get_connection_tienda', lambda: proxy)
    yield {'admin': _admin(), 'cur': cur, 'ids_prod': ids_prod,
           'totales': totales, 'categoria': categoria}
    proxy.rollback()
    raw.close()


@pytest.fixture()
def fx_auth(monkeypatch):
    raw = db_supabase.engine_auth.raw_connection()
    proxy = _Proxy(raw)
    cur = proxy.cursor(dictionary=True)
    # 25 clientes sintéticos: por id DESC quedan al INICIO del listado
    cur.execute(f"""
        INSERT INTO cliente (nombre, apellido, id_tipo_documento,
                             nro_documento, activo, origen)
        SELECT 'testp2-cli-' || gs, 'prueba', 1,
               right(md5(random()::text), 8), TRUE, 'pos'
          FROM generate_series(1, {NSINT}) AS gs
        RETURNING id_cliente
    """)
    ids = sorted(r['id_cliente'] for r in cur.fetchall())
    cur.execute("SELECT COUNT(*) AS n FROM cliente")
    total = cur.fetchone()['n']
    monkeypatch.setattr(modulo_app, 'get_connection_auth', lambda: proxy)
    yield {'admin': _admin(), 'ids': ids, 'total': total, 'cur': cur}
    proxy.rollback()
    raw.close()


# ── /productos (orden: id) ────────────────────────────────────────

def test_productos_pagina_1(fx_tienda):
    total = fx_tienda['totales']['producto']
    html = fx_tienda['admin'].get('/productos').get_data(as_text=True)
    assert f'Mostrando 1–{POR_PAGINA} de {total}' in html
    assert 'aria-current="page"' in html
    # Acciones por fila intactas en página 1
    assert 'data-confirm="¿Eliminar este producto?"' in html
    assert '/productos/editar/' in html
    # Cada sintético está en la página que le toca por posición
    ids = fx_tienda['ids_prod']
    for i, pid in enumerate(ids):
        pos = total - NSINT + 1 + i
        assert (f'>{pid}<' in html) == (_pag_de(pos) == 1), pid


def test_productos_intermedia_y_ultima(fx_tienda):
    total = fx_tienda['totales']['producto']
    ids = fx_tienda['ids_prod']
    # Página que contiene al PRIMER sintético (posición total-24)
    p_med = _pag_de(total - NSINT + 1)
    html = fx_tienda['admin'].get(f'/productos?pagina={p_med}'
                                  ).get_data(as_text=True)
    fin = min(p_med * POR_PAGINA, total)
    assert f'Mostrando {(p_med - 1) * POR_PAGINA + 1}–{fin} de {total}' in html
    assert f'>{ids[0]}<' in html
    # acciones por fila también en páginas > 1
    assert f'/productos/editar/{ids[0]}' in html
    assert 'data-confirm="¿Eliminar este producto?"' in html
    # Última página por clamp: X–Y de Z correcto, el último sintético sí
    # y el primero (posición total-24, siempre en página anterior) no
    ultima = _pag_de(total)
    off = (ultima - 1) * POR_PAGINA
    html = fx_tienda['admin'].get('/productos?pagina=99999'
                                  ).get_data(as_text=True)
    assert f'Mostrando {off + 1}–{total} de {total}' in html
    assert f'>{ids[-1]}<' in html
    assert f'>{ids[0]}<' not in html


def test_productos_pagina_invalida(fx_tienda):
    total = fx_tienda['totales']['producto']
    c = fx_tienda['admin']
    for invalida in ('abc', '0', '-3', '1.5'):
        html = c.get(f'/productos?pagina={invalida}').get_data(as_text=True)
        assert f'Mostrando 1–{POR_PAGINA} de {total}' in html, invalida


def test_productos_filtros_en_href_y_count_filtrado(fx_tienda):
    import html as _htmlmod

    cat = fx_tienda['categoria']
    if not cat:
        pytest.skip('sin categorías reales sin espacios con productos')
    tf = fx_tienda['totales']['categoria']
    crudo = fx_tienda['admin'].get(f'/productos?categoria={cat}&pagina=2'
                                   ).get_data(as_text=True)
    pag = min(2, _pag_de(tf))          # clamp si el filtro cabe en 1 página
    fin = min(pag * POR_PAGINA, tf)
    assert f'Mostrando {(pag - 1) * POR_PAGINA + 1}–{fin} de {tf}' in crudo
    # el COUNT es del conjunto FILTRADO y los href conservan el filtro
    # (el & del atributo viaja HTML-escapado: se des-escapa para asertar)
    assert tf <= fx_tienda['totales']['producto']
    plano = _htmlmod.unescape(crudo)
    assert f'categoria={cat}&pagina=1' in plano


def test_productos_estado_vacio(fx_tienda):
    html = fx_tienda['admin'].get('/productos?categoria=zzz-no-existe'
                                  ).get_data(as_text=True)
    assert 'No se encontraron productos' in html
    assert 'Mostrando' not in html          # macro oculta con total 0


def test_productos_contadores_identicos_en_toda_pagina(fx_tienda):
    patron = re.compile(r'<div class="stat-value">([^<]*)</div>')
    r1 = fx_tienda['admin'].get('/productos').get_data(as_text=True)
    r2 = fx_tienda['admin'].get('/productos?pagina=2').get_data(as_text=True)
    kpis1, kpis2 = patron.findall(r1), patron.findall(r2)
    assert kpis1 and kpis1 == kpis2


def test_productos_export_no_se_pagina(fx_tienda):
    r = fx_tienda['admin'].get('/reportes/productos/exportar?pagina=2')
    assert r.status_code == 200
    filas = _filas_excel(r)
    # cabecera + TODO el conjunto (los 25 sintéticos incluidos)
    assert len(filas) - 1 == fx_tienda['totales']['producto']
    assert sum(1 for f in filas
               if str(f[1] or '').startswith('testp2-prod-')) == NSINT


# ── /inventario (orden: fecha DESC, id DESC) ──────────────────────

def test_inventario_pagina_1(fx_tienda):
    total = fx_tienda['totales']['inventario_movimiento']
    html = fx_tienda['admin'].get('/inventario').get_data(as_text=True)
    assert f'Mostrando 1–{POR_PAGINA} de {total}' in html
    # los sintéticos (fechas recientes) van primero: 20 en página 1
    assert html.count('testp2-mov-') == POR_PAGINA
    # el total del historial es del conjunto completo, no de la página
    assert f'badge bg-secondary ms-2">{total}<' in html


def test_inventario_pagina_2(fx_tienda):
    total = fx_tienda['totales']['inventario_movimiento']
    html = fx_tienda['admin'].get('/inventario?pagina=2').get_data(as_text=True)
    fin = min(2 * POR_PAGINA, total)
    assert f'Mostrando {POR_PAGINA + 1}–{fin} de {total}' in html
    assert html.count('testp2-mov-') == NSINT - POR_PAGINA   # los 5 restantes


def test_inventario_pagina_invalida_y_clamp(fx_tienda):
    total = fx_tienda['totales']['inventario_movimiento']
    c = fx_tienda['admin']
    for invalida in ('abc', '0', '-3', '1.5'):
        html = c.get(f'/inventario?pagina={invalida}').get_data(as_text=True)
        assert f'Mostrando 1–{POR_PAGINA} de {total}' in html, invalida
    ultima = _pag_de(total)
    off = (ultima - 1) * POR_PAGINA
    html = c.get('/inventario?pagina=99999').get_data(as_text=True)
    assert f'Mostrando {off + 1}–{total} de {total}' in html
    assert f'badge bg-secondary ms-2">{total}<' in html


def test_inventario_estado_vacio_en_plantilla():
    fuente = (_TEMPLATES / 'inventario.html').read_text(encoding='utf-8')
    assert 'No hay movimientos registrados' in fuente


def test_inventario_export_no_se_pagina(fx_tienda):
    r = fx_tienda['admin'].get('/reportes/inventario/exportar?pagina=2')
    assert r.status_code == 200
    filas = _filas_excel(r)
    assert len(filas) - 1 == fx_tienda['totales']['inventario_movimiento']
    assert sum(1 for f in filas
               if str(f[6] or '').startswith('testp2-mov-')) == NSINT


# ── /clientes (orden: id DESC) ────────────────────────────────────

def test_clientes_pagina_1(fx_auth):
    total = fx_auth['total']
    html = fx_auth['admin'].get('/clientes').get_data(as_text=True)
    assert f'Mostrando 1–{POR_PAGINA} de {total}' in html
    assert f'Total: {total}' in html          # del conjunto, no de la página
    ids = fx_auth['ids']
    for cid in ids[NSINT - POR_PAGINA:]:      # los 20 mayores: página 1
        assert f'>{cid}<' in html
    for cid in ids[:NSINT - POR_PAGINA]:      # los 5 menores: página 2
        assert f'>{cid}<' not in html
    # acciones por fila intactas
    assert 'data-confirm="¿Eliminar este cliente?"' in html
    assert f'/clientes/editar/{ids[-1]}' in html


def test_clientes_pagina_2(fx_auth):
    total = fx_auth['total']
    ids = fx_auth['ids']
    html = fx_auth['admin'].get('/clientes?pagina=2').get_data(as_text=True)
    fin = min(2 * POR_PAGINA, total)
    assert f'Mostrando {POR_PAGINA + 1}–{fin} de {total}' in html
    for cid in ids[:NSINT - POR_PAGINA]:      # los 5 menores caen aquí
        assert f'>{cid}<' in html
    # acciones por fila también en páginas > 1
    assert f'/clientes/editar/{ids[0]}' in html
    assert 'data-confirm="¿Eliminar este cliente?"' in html


def test_clientes_pagina_invalida_y_clamp(fx_auth):
    total = fx_auth['total']
    c = fx_auth['admin']
    for invalida in ('abc', '0', '-3', '1.5'):
        html = c.get(f'/clientes?pagina={invalida}').get_data(as_text=True)
        assert f'Mostrando 1–{POR_PAGINA} de {total}' in html, invalida
    ultima = _pag_de(total)
    off = (ultima - 1) * POR_PAGINA
    html = c.get('/clientes?pagina=99999').get_data(as_text=True)
    assert f'Mostrando {off + 1}–{total} de {total}' in html


def test_clientes_estado_vacio_en_plantilla():
    fuente = (_TEMPLATES / 'clientes.html').read_text(encoding='utf-8')
    assert 'No hay clientes registrados' in fuente


def test_clientes_export_no_se_pagina(fx_auth):
    r = fx_auth['admin'].get('/reportes/clientes/exportar?pagina=2')
    assert r.status_code == 200
    filas = _filas_excel(r)
    assert len(filas) - 1 == fx_auth['total']
    assert sum(1 for f in filas
               if str(f[1] or '').startswith('testp2-cli-')) == NSINT


# ── /proveedores (orden: nombre, id) ──────────────────────────────

def test_proveedores_pagina_1(fx_tienda):
    total = fx_tienda['totales']['proveedor']
    html = fx_tienda['admin'].get('/proveedores').get_data(as_text=True)
    assert f'Mostrando 1–{POR_PAGINA} de {total}' in html
    assert f'Total: {total}' in html
    # grilla de cards con la paginación debajo y acciones por card
    assert 'class="col-md-4"' in html
    assert 'aria-label="Paginación de resultados"' in html
    assert 'data-confirm="¿Confirma eliminar este proveedor?"' in html


def test_proveedores_ultima_pagina(fx_tienda):
    total = fx_tienda['totales']['proveedor']
    ultima = _pag_de(total)
    off = (ultima - 1) * POR_PAGINA
    html = fx_tienda['admin'].get('/proveedores?pagina=99999'
                                  ).get_data(as_text=True)
    assert f'Mostrando {off + 1}–{total} de {total}' in html
    # el 'zzz-testp2-prov-25' es el último del orden por nombre
    assert 'zzz-testp2-prov-25' in html
    # acciones por card también en la última página
    assert '/proveedores/editar/' in html
    assert 'data-confirm="¿Confirma eliminar este proveedor?"' in html


def test_proveedores_pagina_invalida(fx_tienda):
    total = fx_tienda['totales']['proveedor']
    c = fx_tienda['admin']
    for invalida in ('abc', '0', '-3', '1.5'):
        html = c.get(f'/proveedores?pagina={invalida}').get_data(as_text=True)
        assert f'Mostrando 1–{POR_PAGINA} de {total}' in html, invalida


def test_proveedores_estado_vacio_en_plantilla():
    fuente = (_TEMPLATES / 'proveedores.html').read_text(encoding='utf-8')
    assert 'No hay proveedores registrados' in fuente


def test_proveedores_export_no_se_pagina(fx_tienda):
    r = fx_tienda['admin'].get('/reportes/proveedores/exportar?pagina=2')
    assert r.status_code == 200
    filas = _filas_excel(r)
    assert len(filas) - 1 == fx_tienda['totales']['proveedor']
    assert sum(1 for f in filas
               if str(f[1] or '').startswith('zzz-testp2-prov-')) == NSINT
